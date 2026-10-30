"""Build the short 1.15 Word patch summary after final review.

Use managed-runtime python-docx. Generation still requires rendering and visual
review before delivery; source tests do not attest signed/hardware deployment.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/releases/Angerona_1.15_Security_Patch_Summary.docx"
STAMP = datetime(2026, 10, 30, tzinfo=timezone.utc)


def paragraph(doc, text, *, lead=None):
    result = doc.add_paragraph()
    if lead:
        result.add_run(lead + " ").bold = True
    result.add_run(text)
    return result


def link(doc, title, url):
    result = doc.add_paragraph()
    result.paragraph_format.space_after = Pt(3)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), doc.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
    run, properties = OxmlElement("w:r"), OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "175A8E")
    properties.append(color)
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = title
    run.append(value)
    hyperlink.append(run)
    result._p.append(hyperlink)


def residual_page(doc, title, sections, references):
    doc.add_page_break()
    doc.add_heading("Residual exposure and next steps", 1)
    doc.add_heading(title, 2)
    for label, text in sections:
        paragraph(doc, text, lead=label + ":")
    for name, url in references:
        link(doc, name, url)


def residual_matrix(doc):
    doc.add_page_break()
    doc.add_heading("Cross cutting exposure and priorities", 1)
    paragraph(doc, "Priority reflects the next engineering or deployment dependency, not a prediction that a compromise occurred. Configuration and lifecycle receipts stay separate from detector efficacy and independently administered trust.")
    rows = [
        ("Native availability", "Logging arming and retained evidence", "A native PySide shutdown access violation remains unresolved. One later exit zero does not identify or repair its cause.", "P0"),
        ("Same-user source", "Preserved edits; scoped keys and IPC", "A same-SID attacker can modify writable source, settings and unsealed evidence. Source anchors are not independent trust.", "P1"),
        ("Kernel trust", "Provider-specific canaries; explicit gaps", "A hostile kernel can suppress or forge observations. A green userspace heartbeat cannot attest kernel truth.", "P1"),
        ("Peer suspension", "WER precursor correlation; EKG deadlines", "Simultaneous user-mode peer freeze can stop both detection and recovery. No certified PPL boundary is deployed.", "P1"),
        ("Driver policy", "Enrolled audit/approve/enforce workflow", "Unsigned policy can be removed by an administrator; loaded drivers, boot fallback and reboot effects need acceptance.", "P1"),
        ("Network egress", "Exact-target transaction custody", "Already open or inherited sockets and trusted hosts remain outside path-rule proof; no PID-bound ALE/flow teardown.", "P1"),
        ("Independent history", "Remote signatures and fresh head checks", "Authority compromise or whole-root rollback, unsealed intervals and archive exhaustion remain separate risks.", "P1"),
        ("Coverage and load", "Visible loss/state and bounded work", "Unavailable providers, heuristic false positives, CPU cost and delayed UI can reduce review effectiveness.", "P2"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = (1.05, 1.65, 3.75, .55)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, value in zip(table.rows[0].cells, ("Boundary", "Implemented control", "Remaining exposure", "Next")):
        cell.text = value
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                element = OxmlElement("w:" + edge)
                element.set(qn("w:w"), "90")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right"):
                element = OxmlElement("w:" + edge)
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "4")
                element.set(qn("w:color"), "D9D9D9")
                borders.append(element)
            properties.append(borders)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "ECECEC")
                properties.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                if column_index == 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = index == 0
    paragraph(doc, "P0: diagnose the native shutdown callback/ownership fault. A saved dump includes QObject/widget destruction and event dispatch but no exact root cause. A later normal native exit and an inert negative reproduction do not prove it fixed. The COM exception logged while events continued was not the exit failure. Separately review unattributed diagnostic-log disappearance and durable capture.")
    paragraph(doc, "P1: finish independent/signed deployment and disposable-host acceptance before claiming active protection for these boundaries. P2: profile Core/inventory CPU, improve bounded WER-cache diagnostics and validate longer realistic workloads without relaxing identity or policy.")
    paragraph(doc, "Coverage includes gaps, not just positive detections. The reviewed host had native ETW/Sysmon unavailable, HVCI exposure and held listener attestation; none proves an attack. RWX/JIT, timing and tool matches need corroboration. Model text, imported reports and AI review annotations never grant response authority.")
    link(doc, "Sanitized validation, capture gaps and remaining work",
        "https://github.com/Ag3nt47/AngeronaSuite/blob/main/analysis/hardening-v1.15-2026-10-30.md")


def residual_appendix(doc):
    residual_page(doc, "Kernel telemetry blinding and vulnerable drivers", [
        ("Attacker prerequisites", "A vulnerable-driver attack generally needs a path to load a kernel driver or abuse an already loaded vulnerable component. Administrative or kernel execution puts the attacker beyond the ordinary source process boundary. A valid vendor signature does not by itself make a driver safe."),
        ("Implemented mitigation", "Canaries distinguish actual ETW/Sysmon provenance from process polling and userspace pipeline echoes. A nonce-bound benign process challenge has retained native identity, bounded deadlines and recorder acknowledgment. The optional v2 driver-policy service admits fixed reviewed materials, independent approvals and exact native configuration outcomes."),
        ("What remains possible", "An attacker controlling the kernel can remove callbacks, interfere with delivery or fabricate inputs while user-mode peers stay alive. A forged matching event may defeat a software-only loopback; the canary is not independent kernel attestation. The current driver profile is unsigned/removable and may be changed by an administrator. It does not unload an active driver or restore stripped callbacks."),
        ("Detector limits", "A missed echo establishes a scoped availability problem with unconfirmed cause. A provider that never became available cannot prove new blinding. The observed userspace completions did not exercise a successful native ETW/Sysmon loopback. Fresh EKG liveness and zero queue drops cannot close that gap."),
        ("Deployment needed", "Trusted signed installation, protected service identity/materials, native rights, supported ConfigCI/CiTool, separately reviewed audit history and enforce/reboot/removal recovery acceptance. Microsoft blocklist/HVCI controls require their actual supported OS configuration; this source update does not enable them."),
        ("Validation versus modeling", "Focused tests exercise contracts, provenance, native-owner failure paths and service composition. No vulnerable driver was loaded, no policy was activated, no kernel callbacks were challenged and no signed clean-VM enforcement/reboot acceptance was performed."),
        ("Priority next steps", "P1: validate the signed installed policy lifecycle in disposable supported Windows builds, then verify a genuine native-provider round trip and loss/restart handling. Preserve kernel_prevention_verified:false until a separately defined prevention claim has evidence."),
    ], [("Microsoft recommended driver block rules", "https://learn.microsoft.com/en-us/windows/security/application-security/application-control/app-control-for-business/design/microsoft-recommended-driver-block-rules")])

    residual_page(doc, "WER precursors and user mode peer suspension", [
        ("Attacker prerequisites", "The attacker needs effective access to the target process or an allowed dump/suspension path. A local administrator or compromised peer may have materially greater process-control access than an ordinary unrelated process. The exact OS protection level and permissions determine what is possible."),
        ("Implemented mitigation", "A bounded fast path correlates creation and access observations with a current cached tuple of native-owned PID, birth, role and image context. Stronger correlation requires authenticated source identities/GUIDs. Watchdog/heartbeat state is tied to advancing exact-generation evidence rather than mere PID presence."),
        ("What remains possible", "Post-event observation can arrive after the vulnerable action. A sufficiently privileged attacker may suspend the Core, Watchdog, Scanner and recovery peers together, preventing local detection or respawn. A frozen Core can leave its last UI state visible until another independent observer detects staleness. No guaranteed pre-dump interceptor is deployed."),
        ("Detector limits", "Dump parameters and suspend-capable access are precursors, not evidence that threads actually stopped. Missing process context remains a coverage gap. The earlier Full observation contained six intermittent owned-cache warnings with an unresolved exact cause; a ready inventory or fresh generic heartbeat is not proof of the cache lease at the relevant call."),
        ("Deployment needed", "Certified ELAM-backed antimalware PPL requires the appropriate signing and native installation prerequisites. The implemented virtual-account service separation is useful custody isolation but is not PPL or WinTCB protection. An independently administered observer is needed for stronger assurance against simultaneous local peer failure."),
        ("Validation versus modeling", "Tests cover bounded queues, exact generation/provenance, lifecycle custody and wording. The benign Full capture showed live peers and no WER precursor or actual suspension finding. It was not a coordinated freeze experiment or proof that a privileged attack would be blocked."),
        ("Priority next steps", "P1: complete the signed service deployment/acceptance path and determine certification feasibility. P2: expose bounded typed WER-cache failure/age diagnostics and test benign provider-loss/freeze detection in a disposable authorized environment, retaining the same identity checks."),
    ], [("Microsoft protected antimalware service requirements", "https://learn.microsoft.com/en-us/windows/win32/services/protecting-anti-malware-services-")])

    residual_page(doc, "Containment replacement races and residual egress", [
        ("Attacker prerequisites", "A malicious process must remain active or control a cooperating process, executable namespace or existing network handle during the response window. Privilege, file sharing and process-handle rights constrain which replacement or socket-transfer paths are possible."),
        ("Implemented mitigation", "The response path holds a native process handle plus executable and ancestor identity leases across durable intent, firewall dispatch, postcondition, completion and compensation. File/path admission refuses reparses, remote ambiguity, unsupported hardlinks and other uncertain identity cases. Unknown native completion/close results retain their owners and fence replacement."),
        ("What remains possible", "Detection still precedes admission; earlier exfiltration cannot be undone. The file held for the transaction is not proof of the original mapped image. A path firewall outcome does not bind each socket to PID and birth. Existing or duplicated/inherited sockets, a cooperating already trusted host process and traffic outside the admitted adapter scope remain distinct gaps."),
        ("Detector limits", "Observed connections and configuration readback do not prove packet drops, complete interface/profile scope or termination of an established flow. A failed/unknown postcondition cannot become a completed containment receipt. Suspending a PID alone is not a complete egress guarantee."),
        ("Deployment needed", "The accepted response service/native rights and exact action policy must be deployed. Stronger process-bound enforcement needs a separately designed and accepted native ALE/WFP admission path and established-flow handling, including IPv4/IPv6, interface changes, inheritance and recovery. The current update does not supply that backend."),
        ("Validation versus modeling", "Focused containment tests exercise identity retention, races, journals, compensation and uncertain cleanup. They do not demonstrate a native PID-bound connect decision, inherited-socket revocation or TCP teardown on the reviewed host. Source Observe kept mutation authority held during observation."),
        ("Priority next steps", "P1: define the native process/socket contract and implement it behind enrolled authority, then validate actual traffic outcomes and rollback in a disposable installation. Keep path-rule and socket-admission receipts separate. Preserve emergency managed-only recovery and avoid blanket firewall resets."),
    ], [("Microsoft CreateFileW sharing and handle semantics", "https://learn.microsoft.com/en-us/windows/win32/api/fileapi/nf-fileapi-createfilew")])

    residual_page(doc, "OS downgrade hardware freshness and whole root replay", [
        ("Attacker prerequisites", "An OS servicing downgrade needs control of an applicable privileged servicing/recovery path, affected OS state or storage. Local evidence replay needs the ability to restore a coherent older data/key state. Authority-root rollback and hardware cloning are separate threat boundaries."),
        ("Implemented mitigation", "The audit parses only bounded recognized direct servicing operations from the registry-selected local queue and compares fixed critical disk version resources. It separates Windows.old recovery copies from scheduled replacement. Independent providers are composed before state admission. TPM code verifies pinned fresh counter/extend/counter certifications bound to a typed checkpoint."),
        ("What remains possible", "Unsupported/custom queues, unavailable files and unobserved servicing paths limit detection. A signed older binary can still be vulnerable. A local matching HMAC/database replay cannot prove its own age. A compromised or rolled-back independent authority can undermine freshness. No automatic physical TPM enrollment or protected NV writer is supplied."),
        ("Detector limits", "VERSIONINFO is a best-effort disk sample, not signer, loaded-image, security-build applicability or prevention proof. TPM PolicyNV is a precheck, not atomic digest compare-and-swap. Split writes or reply loss need reconciliation. A software vTPM or cloned authority cannot be assumed physically independent; trusted time and complete authority-root rollback resistance are not established."),
        ("Deployment needed", "Supported vendor OS mitigations remain necessary. Hardware use needs authenticated physical EK/AK enrollment, exact issuer-pinned NV names/policy, suitable platform-created non-ORDERLY counter/extend indexes and a protected serialized single-use-policy writer. These are external provisioning and acceptance tasks, not dashboard toggles."),
        ("Validation versus modeling", "Offline tests cover parser bounds, registry references, signed certification formats, nonce consumption, typed binding and pure policy plans. No downgrade was staged, no reboot or physical TPM call/write was performed, and no hardware floor was active in the observed source instance."),
        ("Priority next steps", "P1: complete independently reviewed physical enrollment/writer and crash-reconciliation design before runtime binding. Test vendor-supported servicing/recovery and real signed hardware observations in disposable systems. Retain unavailable paths and older recovery artifacts without blocking dashboard startup or silently resetting baselines."),
    ], [("Microsoft guidance for blocking VBS rollback", "https://support.microsoft.com/en-us/servicing/os/windows/2024/07/guidance-for-blocking-rollback-of-virtualization-based-security-vbs-related-security-updates")])

    residual_page(doc, "Same user compromise sealed history and service custody", [
        ("Attacker prerequisites", "A malicious package, extension, infostealer or other process executing as the interactive user can reach data and processes permitted to that SID. A source checkout, its batch launcher and locally writable trust anchors share that boundary. Administrator, service-account and authority compromise are stronger separate cases."),
        ("Implemented mitigation", "Independent recorder signatures commit exact timestamp/module/severity/message/details/HMAC bytes and ordered row commitments. Dedicated remote admission enforces predecessor, next epoch and nonoverlap, even with a stolen host request credential. The signing key stays remote. Installed engine admission checks exact service/token/package/root custody before keys or sensors; its enrolled console is status-only."),
        ("What remains possible", "A same-user attacker may alter unsealed history, future local events, source code or the UI; stealing the request credential may allow new malicious append requests. Seals authenticate committed bytes, not truthful sensors. A compromised signing authority or coherent rollback of its whole trusted root can defeat that independence. Removing local archives creates explicit loss, not recoverable original evidence."),
        ("Detector limits", "The worker checks bounded historical ranges, not all history every cycle. Witnessing defaults off and pauses on unavailability or exhausted storage; the default 128 epochs may fill in about two hours. There is no automatic rollover, Python secure erasure or deployed key-ratchet claim. Same-SID diagnostic access remains inside the user's trust boundary; advisory review cannot authorize mutation."),
        ("Deployment needed", "A separately administered authority outside host snapshot/restore jobs, distinct request/signing keys, pinned private HTTPS, mTLS, fresh signed heads and a remote continuity floor. Authority-root freshness needs its own independent generation floor or accepted hardware. Service custody needs real signed installer/SCM provisioning and credentials under the actual service account."),
        ("Validation versus modeling", "Tests cover exact bytes, stolen-credential append restrictions, missing/changed rows, bounded SQL, pending/reply-loss reconciliation, protected enrollment and STOP custody. Synthetic screenshots exercise no real keys/services. Neither an independent production authority nor isolated signed service was deployed during the Full capture."),
        ("Priority next steps", "P1: accept the independent deployment and protected service credential lifecycle on clean hosts, then define export/retention/rollover without discarding unreviewed history. Test authority backup/restore/revocation and crash recovery. Keep archival authenticity, current freshness and event truth as separate results."),
    ], [("Microsoft service security contexts", "https://learn.microsoft.com/en-us/windows/win32/services/service-user-accounts"),
        ("Independent recorder enrollment and limits", "https://github.com/Ag3nt47/AngeronaSuite/blob/main/docs/recorder-integrity.md")])
    residual_matrix(doc)


def native_metadata_followup(doc):
    """Prepared source only; final observation and render review stay separate."""
    doc.add_page_break()
    doc.add_heading("Windows metadata and validation follow up", 1)
    paragraph(doc, "New recorder pending, scratch and writer-lock files receive explicit current-token-user ownership and a protected private DACL before data. Opened-object security, file identity and single-link checks gate admission. Existing inherited private files are accepted; unsuitable existing objects remain preserved and held.", lead="Implemented:")
    paragraph(doc, "Native-to-CRT transfer has one owner. Uncertain creation, transfer or close retains custody and fences optional recorder replacement while the SQLite protection writer remains usable. Default lease callers and POSIX mode 0600 remain unchanged. Parent-directory checks are path checks, not retained directory leases; same-user, administrator and kernel limitations remain.", lead="Boundary:")
    paragraph(doc, "Full Setup exposes optional public witness enrollment/storage values and points to running-dashboard protected identity controls. It stores TLS file paths, not private key contents. Witnessing stays off by default; its unavailability does not block local protection. Release fixtures use explicit installed versions while product equal-version and downgrade rejection remain intact.", lead="Setup and updates:")
    paragraph(doc, "Recorder checks covered 82 distinct cases across two runs: 79 passes and three fixture failures, then three corrected-fixture passes on unchanged product bytes. This is not a single 82-pass run. Setup/updater checks reported 167 passes. Overlapping local cohorts are not summed, and local repairs are separate from fresh hosted CI. The ordinary probe did not reproduce the proposed elevated-CI token-owner cause.", lead="Validation:")
    paragraph(doc, "The final changed-source 120-second Full/STOP met bounded lifecycle acceptance: native exit zero, all 52 observed generations retired and unchanged source/preferences. Yet 17 workers were critical, 34 degraded, one failed and 69 sensor states unknown. Fresh status and paint followed the single Full invocation after about 60 seconds; this is a settling delay, not click latency. The earlier P0 native crash remains unresolved.", lead="Latest observation:")
    paragraph(doc, "Core and inventory CPU remained substantial; separate native requests had partial ordinary-user coverage and a changing-process profile was held. A batch-map improvement changed the private observer only. No product speedup, complete child discovery or all-detector pass follows. Retained historical diagnostics are not fresh collector verification.", lead="Operational limits:")
    paragraph(doc, "Messages now distinguish assurance gaps, KEV review candidates and configured startup references from proven regressions, applicability and execution. Severity, evidence and approval decisions remain. The reviewed redacted alerts did not establish either compromise or a clean host.", lead="Alert precision:")
    link(doc, "Sanitized follow-up evidence and remaining acceptance",
        "https://github.com/Ag3nt47/AngeronaSuite/blob/main/analysis/hardening-v1.15-2026-10-30-followup.md")


def build(*, service_reviewed=False):
    if not service_reviewed:
        raise ValueError("final service/custody review must complete before producing the summary")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    section.footer_distance = Inches(.25)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.04
    for name, size in (("Title", 24), ("Subtitle", 11), ("Heading 1", 13), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(5)
        for border in list(style.element.get_or_add_pPr().findall(qn("w:pBdr"))):
            border.getparent().remove(border)
    props = doc.core_properties
    props.title = "Angerona 1.15 Security Patch Summary"
    props.subject = "Five exploit classes: implemented mitigations and deployment boundaries"
    props.author = props.last_modified_by = "Angerona maintainers"
    props.created = props.modified = STAMP
    props.revision = 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Angerona 1.15.0 | Security patch summary | ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.add_paragraph("Angerona 1.15 Security Patch Summary", "Title")
    doc.add_paragraph("Fixes and residual exposure with deployment requirements", "Subtitle")
    paragraph(doc, "Version 1.15.0 | 30 October 2026")
    paragraph(doc, "This major update strengthens five user-mode security boundaries. These are implemented source controls, not a claim that every exploit is blocked or that Windows vulnerabilities have been patched. Source Observe, installed authority and independent services remain distinct.")
    paragraph(doc, "A native shutdown access violation remains an unresolved P0 availability finding. The later logging patch improves observability; one subsequent normal exit does not establish a crash fix.", lead="Development status:")

    doc.add_heading("1 Kernel telemetry blinding and vulnerable drivers", 1)
    paragraph(doc, "Provider-specific ETW/Sysmon canaries retain a benign nonce probe, native generation, collector provenance and recorder acknowledgment. Missing echoes and unavailable collectors stay explicit even if the heartbeat is alive. The kernel control tick continues in Chill.", lead="Implemented:")
    paragraph(doc, "Enrolled installed v2 driver policy uses pinned materials, server-owned plans, two independent OS-bound approvals, durable capabilities and native readback. The policy is unsigned and removable; default v1 stays diagnostic-only.", lead="Policy workflow:")
    paragraph(doc, "No accepted host policy, intact-callback proof, HVCI activation or guaranteed driver prevention is claimed.", lead="Boundary:")

    doc.add_heading("2 WER dump and suspension precursors", 1)
    paragraph(doc, "A bounded event path correlates WER creation/access observations with current native-owned PID/birth/role context. Stronger correlation requires source identities and GUID evidence; incomplete records remain leads or coverage gaps.", lead="Implemented:")
    paragraph(doc, "Post-event correlation cannot guarantee pre-dump blocking. User-mode peers share a suspension failure domain. Certified ELAM/PPL protection is not supplied.", lead="Boundary:")

    doc.add_heading("3 Executable replacement during containment", 1)
    paragraph(doc, "OS process, executable and ancestor identity leases remain held through durable intent, firewall dispatch, postconditions, completion and managed compensation. Ambiguous paths and uncertain cleanup are held with native owners retained.", lead="Implemented:")
    paragraph(doc, "The acquired file is not original mapped-image proof. PID-bound ALE admission, inherited-socket control and established-flow teardown require a native backend.", lead="Boundary:")

    doc.add_page_break()
    doc.add_heading("4 OS downgrade and whole root rollback", 1)
    paragraph(doc, "Bounded servicing review follows recognized direct operations in the registry-selected local queue and compares fixed critical disk VERSIONINFO resources. Older recovery copies are separate from scheduled replacements. Independent freshness is bound before persisted-state admission.", lead="Implemented:")
    paragraph(doc, "The TPM adapter verifies issuer-pinned physical enrollment and fresh counter/extend/counter certifications bound to a checkpoint. It builds policy-write plans; it does not enroll a device or write NV. Policy prechecks are not atomic compare-and-swap.", lead="Hardware preparation:")
    paragraph(doc, "Vendor mitigations, physical enrollment and a protected serialized external writer remain required. Disk versions do not prove signer, loaded image or patch applicability. Whole-authority rollback is separate.", lead="Boundary:")

    doc.add_heading("5 Same user tampering and history forgery", 1)
    paragraph(doc, "Optional independent Ed25519 signatures commit exact recorder row bytes. Remote admission enforces the next predecessor/epoch/nonoverlapping range; a durable pending claim reconciles reply loss. Protected request-key enrollment and retained workers preserve custody through STOP.", lead="Implemented:")
    paragraph(doc, "Installed engine admission checks the actual virtual account, SCM process, restricted token, package and protected root before keys or sensors. Its console is status-only. Signed installer provisioning and acceptance remain separate.", lead="Service separation:")
    paragraph(doc, "Witnessing defaults off. Its 16 MiB/128-epoch archive preserves and pauses when full, potentially after roughly two hours at a 60-second cadence. Source Observe shares its user's failure domain. No Python secure-erasure or certified PPL claim is made.", lead="Boundary:")

    doc.add_heading("Validation and operational findings", 1)
    paragraph(doc, "Focused cohorts: native/service/status 183 passes; recorder 64; credential composition 71 plus one POSIX skip; TPM 83; retrieval 11; watchdog nine; socket presentation 39; crash-log observability 55. Cohorts overlap and are not summed. Broad selfcheck: 28 phase passes and one existing host Ollama listener-attestation failure; its 30-step defensive drill passes.")
    paragraph(doc, "The earlier 180-second Full observation retained a 4,446-record initial gap. A later logging-enabled 120-second session retained a separate 2,898-record gap and 268 new events, with 75/75 workers running but native ETW/Sysmon unavailable and 69 collector states unknown. No confirmed breach was established. CPU cost and UI delays remain; the sessions are not a controlled benchmark.")
    paragraph(doc, "That session's STOP retired all retained generations with actual Core exit zero and unchanged source/preferences. Observation acceptance still failed on two process-inventory bounds; count versus deadline cause is unknown. A prior session's native shutdown crash remains unresolved. Both crash logs later disappeared for an unproved reason; private diagnostic text is preserved. The 55-pass logging patch verifies arming visibility, not complete capture or crash repair. The subsequent follow-up records the final changed-source lifecycle result separately.")
    link(doc, "Release scope, evidence and deployment limits",
        "https://github.com/Ag3nt47/AngeronaSuite/blob/main/docs/releases/v1.15.0.md")
    link(doc, "Master Manual security flow",
        "https://github.com/Ag3nt47/AngeronaSuite/blob/main/docs/security-flow.md")
    native_metadata_followup(doc)
    residual_appendix(doc)
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--service-reviewed", action="store_true",
                        help="Required after final service/custody review; render QA is still separate")
    args = parser.parse_args()
    data = BytesIO()
    build(service_reviewed=args.service_reviewed).save(data)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(data) as source, ZipFile(args.output, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2026, 10, 30, 0, 0, 0))
            info.compress_type, info.create_system = ZIP_DEFLATED, 0
            target.writestr(info, source.read(name))
    print(f"{args.output} (render every page before delivery)")


if __name__ == "__main__":
    main()
