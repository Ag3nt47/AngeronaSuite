"""Append the October 9 operator update while preserving historical chapters."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import shutil

from docx import Document
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "Angerona_Master_Manual.docx"
BASELINE = ROOT / ".tmp/manual_simulation_oct09/Angerona_Master_Manual_before.docx"
TITLE = "Operator update 9 October 2026"


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def update():
    BASELINE.parent.mkdir(parents=True, exist_ok=True)
    if not BASELINE.exists():
        if any(p.text == TITLE for p in Document(MANUAL).paragraphs):
            raise ValueError("Update already applied; use the retained baseline to rebuild")
        shutil.copy2(MANUAL, BASELINE)
    document = Document(BASELINE)
    for paragraph in document.paragraphs:
        if paragraph.text == "8 October 2026":
            replace_text(paragraph, "9 October 2026")
    for row in document.tables[1].rows:
        if row.cells[0].text in {"Date", "Updated", "Last updated"}:
            replace_text(row.cells[1].paragraphs[0], "9 October 2026")
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("8 October 2026", "9 October 2026")
    heading = document.add_paragraph(TITLE, "Heading 1")
    heading.paragraph_format.page_break_before = True
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    paragraphs = [
        "Simulation preparation and stop cleanup now run outside the dashboard event loop. The console shows preparation while target checks, detector readiness and engine startup complete. Keep the console open to see the handoff. Stop clean cancels pending preparation and retires the run it owns; a cleanup failure remains visible. Duplicate launches wait for preparation or cleanup to finish. Selected exercises and response rules retain their existing behavior.",
        "The October 9 comprehensive campaign recorded 27 of 80 planned mandatory steps before a validation-producer liveness failure. Its manifest is incomplete and ineligible for scoring. Executed steps are not detections or containment receipts. No finalized October 9 after-action report was available during review; the September 23 report is historical. Live producer and signed-evidence checks remain required, and the failed producer has not been identified from the generic error.",
        "Repeated BYOVD drill phases now use distinct inert marker paths. This prevents a marker from an earlier phase blocking a later exclusive creation. Detector matching accepts the exact generated name, while response eligibility still requires the registered file identity and content. A similar filename alone grants no authority. Worker-start failure also releases Shark's run state so a later attempt can start.",
        "Black Box refreshes firewall and SOAR snapshots in background workers and retains the prior view if a refresh fails. SOAR shows at most 200 valid records from the latest 1 MiB. Firewall output, retained rules and visible rows have explicit limits. Log tails drain bounded chunks without deleting or skipping source evidence; large crash logs show a labeled preview. CPU readings now retain a sample for the same process identity.",
        "Scanners cooperate with the existing responsiveness governor during directory enumeration and file batches. Cancellation remains interruptible, deadlines still apply and partial inventories do not become completed coverage. The standalone scanner closes its resources on failure and cannot start a second sensor loop after a GUI error. Driver inventory has a finite watchdog budget matching its collector timeout; routine network-trust inventory can wait longer under pressure while collecting fresh evidence.",
        "Sandbox Editor now tests the opened buffer, safely retires completed workers and preserves undo entries when a revert fails. Source files are limited to 4 MiB; undo retains at most 20 copies and 16 MiB. Trusted baseline tests have bounded child output, memory and cleanup. Candidate edits remain AST-only. This editor is not an isolation boundary for executing hostile code, and working copies do not replace production sources.",
        "Flow windows cancel the initial AI poll when closing and reject late worker updates. These changes address confirmed blocking and lifetime bugs; they do not guarantee a maximum response time under Windows load. Restart Angerona and Black Box to load the update. See analysis/simulation-responsiveness-2026-10-09.md for the review and validation. GitHub Actions identifies the exact published commit; earlier manual results retain their historical dates.",
    ]
    for text in paragraphs:
        document.add_paragraph(text, "Normal")
    document.core_properties.modified = datetime(2026, 10, 9, tzinfo=timezone.utc)
    document.core_properties.subject = "Operator reference with simulation responsiveness and bounded diagnostic work"
    document.save(MANUAL)
    assert sum(p.text == TITLE for p in Document(MANUAL).paragraphs) == 1
    print(f"Updated {MANUAL.name}")


if __name__ == "__main__":
    update()
