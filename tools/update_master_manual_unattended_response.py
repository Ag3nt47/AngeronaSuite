"""Append the October 8 operator update while preserving historical chapters."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import shutil

from docx import Document
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "Angerona_Master_Manual.docx"
BASELINE = ROOT / ".tmp/manual_unattended_response/Angerona_Master_Manual_before.docx"
TITLE = "Operator update 8 October 2026"


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def update():
    BASELINE.parent.mkdir(parents=True, exist_ok=True)
    if not BASELINE.exists():
        if any(p.text == TITLE for p in Document(MANUAL).paragraphs):
            raise ValueError("Update already applied; use the retained baseline to rebuild")
        shutil.copy2(MANUAL, BASELINE)
    document = Document(BASELINE)
    for paragraph in document.paragraphs:
        if paragraph.text == "4 October 2026":
            replace_text(paragraph, "8 October 2026")
    for row in document.tables[1].rows:
        if row.cells[0].text in {"Date", "Updated", "Last updated"}:
            replace_text(row.cells[1].paragraphs[0], "8 October 2026")
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("4 October 2026", "8 October 2026")
    heading = document.add_paragraph(TITLE, "Heading 1")
    heading.paragraph_format.page_break_before = True
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    paragraphs = [
        "Benign Shark and Red Team simulations continue when automatic response is unavailable. An amber warning remains visible through progress and results. Selected profiles, intensity and inert process/file exercises still run. Target ownership, the validation lease and engine startup must still succeed. Only authenticated verified response receipts count as containment; a completed exercise with no receipts remains unverified.",
        "Adversary Combat already defaults to automatic standing authority. Eligible detector evidence invokes the saved file, process, network, host-isolation and deception rules without asking for each incident and without requiring Ollama. The policy is restricted to exact authenticated detector-authorized targets. Health findings, exposure and model advice do not become containment commands.",
        "The footer shows Auto ARMED, HELD, OFF, STARTING or UNAVAILABLE beside posture. Settings > Adversary Combat describes the effective rules and retains readiness, verified action history and Undo. Contain mode suspends eligible processes even if a different process action is saved. An armed policy cannot guarantee that every alert has an eligible target or that every action succeeds.",
        "A narrowly defined interrupted simulation checkpoint can now recover automatically during enabled Combat startup. Recovery requires an authenticated journal with exactly one trailing quarantine intent, agreeing protected checkpoint and independent witness, and no ambiguous pending recovery. The original single-link marker must retain its recorded identity, hash and exact shipped comprehensive-drill bytes in the suite drill sandbox; its quarantine destination must be absent.",
        "Recovery preserves private, verified copies of the unchanged journal, protected store and witness before advancing the matching checkpoint. It neither moves the marker nor creates signing authority, truncates history or reports a quarantine success. Normal reconciliation records the unapplied intent. Changed files, real containment effects, unrelated files, unsupported history or an unsuccessful proof leave automatic response held for verified recovery.",
        "When restarting this Windows source installation, use kill-all-angerona.bat with elevation to stop the verified Angerona processes, then reopen start-angerona.bat in the normal user session. The source launcher intentionally runs unelevated in Observe/development mode; the separately installed signed distribution supplies Windows Protect boundaries. Do not delete recovery data to clear a warning.",
        "The implementation record, regression evidence and remaining limits are in analysis/unattended-response-2026-10-08.md. Earlier validation totals in this manual retain their historical dates. GitHub Actions and the guarded publisher establish the exact published commit and current CI outcome.",
    ]
    for text in paragraphs:
        document.add_paragraph(text, "Normal")
    document.core_properties.modified = datetime(2026, 10, 8, tzinfo=timezone.utc)
    document.core_properties.subject = "Operator reference with simulation warnings and unattended response recovery"
    document.save(MANUAL)
    assert sum(p.text == TITLE for p in Document(MANUAL).paragraphs) == 1
    print(f"Updated {MANUAL.name}")


if __name__ == "__main__":
    update()
