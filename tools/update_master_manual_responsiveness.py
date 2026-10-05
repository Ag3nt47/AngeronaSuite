"""Apply the October 2026 operator update without rebuilding historical chapters."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path
import shutil

from docx import Document
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "Angerona_Master_Manual.docx"
BASELINE = ROOT / ".tmp/manual_responsiveness/Angerona_Master_Manual_before.docx"
MARKER = "17.10 Dashboard responsiveness and module review October 2026"


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_section(document, anchor_prefix, heading, paragraphs):
    anchor = next(p for p in document.paragraphs
                  if p.style.name == "Heading 1" and p.text.startswith(anchor_prefix))
    for text, style in [(heading, "Heading 2"), *[(p, "Normal") for p in paragraphs]]:
        paragraph = document.add_paragraph(text, style=style)
        if style == "Heading 2" and heading.startswith("5.4 "):
            paragraph.paragraph_format.page_break_before = True
        anchor._p.addprevious(paragraph._p)


def update(validation):
    BASELINE.parent.mkdir(parents=True, exist_ok=True)
    if not BASELINE.exists():
        if any(p.text == MARKER for p in Document(MANUAL).paragraphs):
            raise ValueError("Update already applied; original baseline is required to rebuild")
        shutil.copy2(MANUAL, BASELINE)
    document = Document(BASELINE)
    for paragraph in document.paragraphs:
        if paragraph.text == "30 August 2026":
            set_text(paragraph, "4 October 2026")
        elif paragraph.text == "ANGERONA":
            paragraph.style = document.styles["Title"]
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.underline = False
        elif paragraph.text.startswith("The six pytest skips"):
            set_text(paragraph, "The historical Cycle 25 suite had six declared host/platform capability skips; later gates have their own explicit skip counts. Unknown failures cannot be waived as platform skips.")
        elif paragraph.text.startswith("R6-03 remains open"):
            set_text(paragraph, "R6-03 remains a defense-in-depth proposal: retaining an OS process handle and a bounded executable identity lease across the complete program-firewall action would strengthen identity custody. October review separately reproduced and fixed failed compensation being terminalized despite uncertain effects; unverified reversal now remains in recovery. Neither finding establishes complete host-mutation safety.")

    set_text(document.tables[0].cell(0, 0).paragraphs[0],
             "Current release: v1.13.0 with October dashboard responsiveness, default-on adaptive routine scan pacing, and a review of all 87 module source files covering 84 discovered capabilities.")
    for row in document.tables[1].rows:
        if row.cells[0].text == "Release state":
            set_text(row.cells[1].paragraphs[0], "October 2026 maintenance. See section 13.3 for local validation; publication and exact-commit CI remain governed by the canonical publisher and GitHub Actions.")
        elif row.cells[0].text in {"Date", "Updated", "Last updated"}:
            set_text(row.cells[1].paragraphs[0], "4 October 2026")
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("30 August 2026", "4 October 2026")
    # Keep historical validation and red-team verdicts explicitly dated.
    set_text(document.tables[16].cell(0, 1).paragraphs[0], "Historical Cycle 34 result")
    paragraph = document.tables[18].cell(0, 0).paragraphs[0]
    set_text(paragraph, paragraph.text.replace("Final red-team verdict:", "Historical Cycle 34 red-team verdict:"))

    insert_section(document, "6. Detection", "5.4 Dashboard display and readable module ribbons", [
        "Settings > Appearance > Dashboard display selects Standard or the optional Orbital dashboard. Standard remains the default. Orbital presents existing module and posture snapshots, with click-through details and a Workspace tab for normal operations; it adds no telemetry collectors.",
        "The bottom HEALTH and ACTIVITY ribbons keep individual modules in readable tiles. Slow synchronized scrolling, pause, wheel and keyboard navigation let you browse without shrinking every module into the window. Reduced motion and hidden-window handling stop unnecessary animation. Activity remains an estimate, separate from health, assurance and actual host CPU.",
        "The persistent footer shows CPU, RAM, dashboard FPS, Angerona pace, network rates and posture. It reflows into additional rows as space narrows. Hover or keyboard access provides full metric descriptions when a caption is shortened.",
        "FPS is a tiny Qt footer paint heartbeat targeting 30 deliveries per second. Its percentage is the observed rate divided by 30, capped at 100. It measures dashboard responsiveness, not GPU or monitor refresh rate. Hidden or minimized dashboards show a paused state; startup shows measuring until a real sample is available.",
        "Angerona pace is a separate bar showing the requested routine scan budget: 100, 50, 25 or about 13 percent for interval multipliers 1, 2, 4 or 8. For example, CPU 10 percent and Angerona pace 25 percent describe different quantities. Pace is neither a CPU reservation nor measured throughput or a security coverage percentage.",
    ])
    insert_section(document, "11. Red Team", "10.4 Adaptive pacing for dashboard responsiveness", [
        "Adaptive routine scan pacing is enabled by default, including when older settings files do not contain the option. Change it under Settings > Appearance. The controller reuses existing host CPU observations and the small footer heartbeat; it creates no new host sampling thread.",
        "Sustained slow paints, delayed scheduling, a missing visible heartbeat or high CPU progressively lengthen opted-in routine work. Healthy feedback restores one level every eight seconds. Added interval delay is capped at 30 seconds; cooperative inner-loop checkpoints wait at most 50 milliseconds per batch and remain interruptible by stop.",
        "Paced work includes routine process and network inventories, Memory Time-Machine, memory injection scanning, YARA file scanning, FIM full-file scanning and Sysmon process fallback. FIM hash chunks and memory-region walks yield during long passes. Event delivery, streaming telemetry, response execution, watchdogs and direct self-tests keep their own cadence. Other governors combine by the greater requested delay instead of multiplying delays.",
        "Slower routine scans can increase discovery latency. Critical action authority and evidence checks remain mandatory. This cooperative control cannot preempt a blocked native call or promise a CPU quota or a freeze-free host. On a host that keeps up, the requested pace returns to 100 percent.",
    ])
    insert_section(document, "12. Recovery", "11.3 Responsive sandbox editor", [
        "Opening the Red Team console leaves the unvisited Sandbox Editor unloaded. Its first visit, reload, syntax-checked save and confirmed rollback use one background worker. Busy controls prevent overlapping operations; errors preserve the current buffer and offer retry. All writes remain confined to the isolated source working copy. The installed application is not hot-reloaded.",
    ])
    insert_section(document, "14. Performance", "13.3 October module review and regression evidence", [
        "The October review inspected every one of the 87 Python files under src/angerona/modules and the 84 discovered defensive capabilities. Per-file purpose, security boundaries, performance concerns, fixes and remaining limits are recorded in analysis/module-review-2026-10-04. Review is evidence about this source snapshot, not independent certification.",
        "Fixes cover routine telemetry causing priority escalation, malformed cloud verdicts, uncertain Combat compensation and firewall-state verification, remote observation entering local response correlation, Defender live log continuity, false FRZ health, incomplete forensics and network observations, process identity reuse, provenance cycle checks, bounded SOAR retry state, safe Mobile ECO scope and bounded alert authorization state, nonmutating Intel self-tests, and Smart Deception cancellation.",
        validation,
        "Tests and reproductions use isolated data, inert events and mocked host actions where required. This maintenance did not run a live hostile campaign or exercise actual firewall changes, native process suspension, privileged watchdog isolation or a prolonged elevated all-sensor soak. GitHub Actions records the exact published commit; the publisher must prove public main and the local branch match and public README assets match their tracked bytes.",
    ])
    insert_section(document, "15. Security", "14.3 Sustained responsiveness evidence", [
        "An isolated Memory Time-Machine reproduction found that 257 processes with three unchanged observations caused all 771 observations to be forwarded again on every sweep because the 256-PID index continually evicted the next process. The wider bounded index changes repeat sweeps to zero forwarded observations while preserving the aggregate 1,048,576-fingerprint ceiling and retry after failed queue admission. A 320-process fixture similarly changes 960 redundant observations to zero.",
        "Historic local watchdog logs recorded roughly six-second stalls while constructing a details view or reading the simulation editor. Deferring editor file operations removes a confirmed GUI-thread I/O path. These reproductions show avoided work and bounded control behavior; they do not measure a guaranteed whole-host CPU or FPS improvement. Restart the updated application and observe a normal sustained session to assess remaining load.",
    ])
    insert_section(document, "Appendix A", MARKER, [
        "This v1.13.0 maintenance adds readable responsive dashboard controls, the optional orbital overview, a default-on FPS-guided routine scan controller and a distinct Angerona pacing bar. It also reduces repeated Memory Time-Machine processing and makes sandbox editing asynchronous.",
        "The complete module review applies reproducible bug, security and performance fixes while preserving exact response authority and truthful incomplete states. Source reports retain unresolved operational and scaling limits, including extended native host validation, bounded-query byte limits on older SQLite bindings, and retention policy work that must not silently delete recovery evidence.",
        "Primary records: analysis/sustained-responsiveness-2026-10-04.md and analysis/module-review-2026-10-04/README.md. Earlier chapters retain explicitly historical validation results rather than presenting them as this update's test totals.",
    ])
    document.core_properties.title = "Angerona Master Manual"
    document.core_properties.subject = "Operator reference with October 2026 dashboard pacing and module review"
    document.core_properties.modified = datetime(2026, 10, 4, tzinfo=timezone.utc)
    document.save(MANUAL)
    reopened = Document(MANUAL)
    assert sum(p.text == MARKER for p in reopened.paragraphs) == 1
    assert len(reopened.tables) == 23
    print(f"Updated {MANUAL}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--validation", required=True)
    update(parser.parse_args().validation)
